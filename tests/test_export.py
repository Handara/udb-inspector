import asyncio
import csv
import datetime
import json
import sqlite3
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from textual.widgets import Input, Select

from udb_inspector import export
from udb_inspector.database import Database
from udb_inspector.detect import detect
from udb_inspector.tui.app import InspectorApp
from udb_inspector.tui.query import ExportScreen
from udb_inspector.tui.screens import ConfirmScreen

AWKWARD = (
    1,
    None,
    datetime.datetime(2026, 1, 2, 3, 4, tzinfo=datetime.timezone.utc),
    Decimal("1.50"),
    b"\x00\x01",
    {"a": [1]},
    "bell\x07",
)
AWKWARD_COLUMNS = ["id", "id", "at", "price", "blob", "doc", "text"]


@pytest.fixture
def items_db(tmp_path):
    path = tmp_path / "items.sqlite3"
    conn = sqlite3.connect(path)
    conn.execute("CREATE TABLE items (id INTEGER PRIMARY KEY, name TEXT)")
    conn.executemany("INSERT INTO items (name) VALUES (?)", [(f"item {i:03}",) for i in range(1, 121)])
    conn.commit()
    conn.close()
    return path


def test_csv(tmp_path):
    report = export.write("csv", tmp_path / "x.csv", AWKWARD_COLUMNS, [[AWKWARD]])
    assert report == export.ExportReport(1)
    rows = list(csv.reader((tmp_path / "x.csv").open(encoding="utf-8")))
    assert rows == [AWKWARD_COLUMNS, ["1", "", "2026-01-02T03:04:00+00:00", "1.50", "0001", '{"a": [1]}', "bell\x07"]]


def test_json(tmp_path):
    export.write("json", tmp_path / "x.json", AWKWARD_COLUMNS, [[AWKWARD]])
    assert json.loads((tmp_path / "x.json").read_text()) == [
        {"id": 1, "id_2": None, "at": "2026-01-02T03:04:00+00:00", "price": 1.5, "blob": "0001", "doc": {"a": [1]}, "text": "bell\x07"}
    ]
    export.write("json", tmp_path / "empty.json", ["a"], [])
    assert json.loads((tmp_path / "empty.json").read_text()) == []


def test_xlsx(tmp_path):
    export.write("xlsx", tmp_path / "x.xlsx", AWKWARD_COLUMNS, [[AWKWARD]])
    sheet = load_workbook(tmp_path / "x.xlsx").active
    assert [[cell.value for cell in row] for row in sheet.iter_rows()] == [
        AWKWARD_COLUMNS,
        [1, None, "2026-01-02T03:04:00+00:00", 1.5, "0001", '{"a": [1]}', "bell"],  # Excel rejects \x07
    ]
    assert sheet["A1"].font.bold


def test_xlsx_row_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(export, "EXCEL_MAX_ROWS", 2)
    report = export.write("xlsx", tmp_path / "x.xlsx", ["n"], [[(1,), (2,)], [(3,)]])
    assert report.rows == 2 and "1 were left out" in report.note
    assert load_workbook(tmp_path / "x.xlsx").active.max_row == 3  # header + 2


def test_tab_separated():
    assert export.tab_separated(["a", "b"], [(1, "x\ty"), (None, "two\nlines")]) == "a\tb\n1\tx y\n\ttwo lines\n"


def test_stream_query_gets_every_row(items_db, tmp_path):
    db = Database(detect(str(items_db)))
    with db.stream_query("SELECT * FROM items", batch_size=50) as (columns, batches):
        sizes = [len(batch) for batch in batches]
    assert columns == ["id", "name"] and sizes == [50, 50, 20]


# The export dialog


def run(app, scenario, size=(140, 45)):
    async def go():
        async with app.run_test(size=size) as pilot:
            await scenario(pilot)

    asyncio.run(asyncio.wait_for(go(), 20))


async def settle(pilot):
    for _ in range(4):
        await pilot.app.workers.wait_for_complete()
        await pilot.pause()


async def query_items(app, pilot):
    await settle(pilot)
    await pilot.press("down", "s")
    await settle(pilot)


def test_export_all_rows_runs_the_query_again(items_db, tmp_path):
    app = InspectorApp(detect(str(items_db)), cwd=tmp_path)

    async def scenario(pilot):
        await query_items(app, pilot)
        assert app.query_panel.paged and len(app.query_panel.result.rows) == 50  # one page
        await pilot.press("ctrl+s")
        screen = app.screen
        assert isinstance(screen, ExportScreen)
        assert screen.query_one("#file", Input).value.startswith("items-")
        screen.query_one("#format", Select).value = "xlsx"
        await pilot.pause()
        assert screen.query_one("#file", Input).value.endswith(".xlsx")
        screen.query_one("#file", Input).value = "all.xlsx"
        await pilot.press("ctrl+s")
        await settle(pilot)
        assert "Exported 120 rows to" in str(app.query_one("#status").content)

    run(app, scenario)
    assert load_workbook(tmp_path / "all.xlsx").active.max_row == 121


def test_export_the_rows_shown(items_db, tmp_path):
    app = InspectorApp(detect(str(items_db)), cwd=tmp_path)

    async def scenario(pilot):
        await query_items(app, pilot)
        app.query_panel.query_one("#search", Input).value = "item 11"
        await pilot.pause(0.6)  # the search runs once typing pauses
        await settle(pilot)
        app.query_panel.grid.focus()
        await pilot.press("x")
        screen = app.screen
        assert screen.query_one("#row-scope").display  # searching makes "rows shown" a real choice
        screen.query_one("#scope", Select).value = "view"
        screen.query_one("#file", Input).value = "shown.csv"
        await pilot.press("ctrl+s")
        await settle(pilot)

    run(app, scenario)
    rows = list(csv.reader((tmp_path / "shown.csv").open()))
    assert len(rows) == 11 and rows[1] == ["110", "item 110"]


def test_copy_to_clipboard(items_db, tmp_path):
    app = InspectorApp(detect(str(items_db)), cwd=tmp_path)

    async def scenario(pilot):
        await query_items(app, pilot)
        await pilot.press("ctrl+s")
        app.screen.query_one("#format", Select).value = "clipboard"
        await pilot.pause()
        assert not app.screen.query_one("#row-file").display
        await pilot.press("ctrl+s")
        await pilot.pause()
        assert app.clipboard.startswith("id\tname\n1\titem 001\n")
        assert "Copied 120 rows" in str(app.query_one("#status").content)

    run(app, scenario)


def test_replacing_a_file_asks_first(items_db, tmp_path):
    (tmp_path / "out.csv").write_text("old")
    app = InspectorApp(detect(str(items_db)), cwd=tmp_path)

    async def scenario(pilot):
        await query_items(app, pilot)
        await pilot.press("ctrl+s")
        app.screen.query_one("#file", Input).value = "out.csv"
        await pilot.press("ctrl+s")
        await pilot.pause()
        assert isinstance(app.screen, ConfirmScreen)
        await pilot.press("y")
        await settle(pilot)

    run(app, scenario)
    assert (tmp_path / "out.csv").read_text().startswith("id,name")


def test_nothing_to_export_before_a_query(items_db):
    app = InspectorApp(detect(str(items_db)))

    async def scenario(pilot):
        await settle(pilot)
        assert app.query_panel.check_action("export", ()) is False

    run(app, scenario)
