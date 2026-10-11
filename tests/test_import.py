import asyncio
import datetime
import sqlite3
import uuid
from dataclasses import replace
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import types as t
from sqlalchemy.dialects.oracle import NUMBER
from textual.widgets import Button, Select

from udb_inspector import database
from udb_inspector.database import Database, Relation, WritesNotAllowed
from udb_inspector.detect import detect
from udb_inspector.importing import (
    ColumnInfo,
    convert,
    describe_column,
    match_columns,
    parse_pasted,
    plan,
    save_template,
    template_columns,
    template_tsv,
)
from udb_inspector.tui.app import InspectorApp
from udb_inspector.tui.screens import ConfirmScreen


def column(name, type_, nullable=True, has_default=False, primary_key=False, generated=False):
    return ColumnInfo(name, type_, str(type_), nullable, has_default, primary_key, generated)


COLUMNS = [
    column("id", t.Integer(), nullable=False, has_default=True, primary_key=True, generated=True),
    column("name", t.String(50), nullable=False),
    column("price", t.Numeric(10, 2)),
    column("active", t.Boolean(), nullable=False, has_default=True),
]


# Reading what was pasted


def test_parse_excel_cells():
    pasted = parse_pasted('Name\tNotes\r\nAda\t"two\nlines"\r\n\t\r\nBob\t"has ""quotes"""\r\n\r\n')
    assert pasted.headers == ["Name", "Notes"]
    assert pasted.rows == [["Ada", "two\nlines"], ["Bob", 'has "quotes"']]  # blank rows dropped


@pytest.mark.parametrize("text", ["id,name\n1,\"a, b\"\n", "id;name\n1;a, b\n"])
def test_parse_csv(text):
    assert parse_pasted(text).rows == [["1", "a, b"]]


def test_match_columns_ignores_case_spaces_and_underscores():
    mapping, ignored = match_columns(["NAME", "Price ", "is-active", "Colour", "name"], COLUMNS + [column("is_active", t.Boolean())])
    assert mapping == {0: "name", 1: "price", 2: "is_active"}
    assert ignored == ["Colour", "name"]  # the second "name" is a duplicate


# Converting values


@pytest.mark.parametrize(
    ("type_", "raw", "expected"),
    [
        (t.Integer(), "42", 42),
        (t.Integer(), "1,234", 1234),
        (t.Numeric(10, 2), "1,234.50", Decimal("1234.50")),
        (t.Float(), "2.5", 2.5),
        (t.Boolean(), "TRUE", True),
        (t.Boolean(), "no", False),
        (t.Date(), "2026-10-11", datetime.date(2026, 10, 11)),
        (t.Date(), "2026/10/11", datetime.date(2026, 10, 11)),
        (t.DateTime(), "2026/10/11 9:05", datetime.datetime(2026, 10, 11, 9, 5)),
        (t.DateTime(), "2026-10-11T09:05:30", datetime.datetime(2026, 10, 11, 9, 5, 30)),
        (t.Time(), "09:30", datetime.time(9, 30)),
        (t.JSON(), '{"a": [1]}', {"a": [1]}),
        (t.LargeBinary(), "0a1b", b"\x0a\x1b"),
        (t.Uuid(), "12345678-1234-5678-1234-567812345678", uuid.UUID("12345678-1234-5678-1234-567812345678")),
        (t.String(), "  kept as typed", "  kept as typed"),
    ],
)
def test_convert(type_, raw, expected):
    assert convert(raw, column("c", type_)) == expected


@pytest.mark.parametrize(
    ("type_", "raw", "message"),
    [
        (t.Integer(), "1.5", "isn't a whole number"),
        (t.Numeric(), "cheap", "isn't a number"),
        (t.Boolean(), "maybe", "isn't true or false"),
        (t.Date(), "11/10/2026", "write dates as YYYY-MM-DD"),
        (t.Date(), "2026-10-11 14:00", "has a time"),
        (t.JSON(), "{nope", "isn't valid JSON"),
    ],
)
def test_convert_problems(type_, raw, message):
    with pytest.raises(ValueError, match=message):
        convert(raw, column("c", type_))


def test_oracle_numbers():
    """SQLAlchemy calls every Oracle NUMBER an Integer too; the scale says whether decimals fit."""
    assert convert("1,234.50", column("c", NUMBER(10, 2))) == Decimal("1234.50")
    assert convert("1.5", column("c", NUMBER())) == Decimal("1.5")
    assert convert("7", column("c", NUMBER(10))) == 7
    with pytest.raises(ValueError, match="isn't a whole number"):
        convert("1.5", column("c", NUMBER(10, 0)))


# Planning


def test_plan_insert():
    pasted = parse_pasted("name\tprice\tactive\nAda\t1.5\t\n\t2\tyes\nBob\tNULL\tmaybe\n")
    result = plan(pasted, COLUMNS, "insert", ["id"])
    assert result.problems == []
    ada, nameless, bob = result.rows
    assert (ada.line, ada.values) == (2, {"name": "Ada", "price": Decimal("1.5")})  # active left to its default
    assert (nameless.line, nameless.problem) == (3, "name is required")
    assert bob.line == 4 and "active: 'maybe' isn't true or false" in bob.problem


def test_plan_insert_needs_required_columns_pasted():
    result = plan(parse_pasted("price\n1\n"), COLUMNS, "insert", ["id"])
    assert result.rows[0].problem == "name is required"


@pytest.mark.parametrize(
    ("operation", "text", "problem"),
    [
        ("update", "name\nAda\n", "paste a column for id"),
        ("update", "id\n1\n", "at least one column to change"),
        ("insert", "colour\nred\n", "None of the pasted headers"),
    ],
)
def test_plan_problems(operation, text, problem):
    assert problem in " ".join(plan(parse_pasted(text), COLUMNS, operation, ["id"]).problems)


def test_plan_delete_uses_only_the_key_and_needs_it():
    result = plan(parse_pasted("id\tname\n1\tAda\n\tBob\n"), COLUMNS, "delete", ["id"])
    assert result.mapping == {0: "id"}
    assert result.rows[0].values == {"id": 1}
    assert result.rows[1].problem == "id is empty"


# Templates


def test_template_columns_follow_the_operation():
    names = lambda columns: [c.name for c in columns]  # noqa: E731
    assert names(template_columns(COLUMNS, "insert", ["id"])) == ["name", "price", "active"]  # no generated id
    assert names(template_columns(COLUMNS, "update", ["name"])) == ["name", "id", "price", "active"]  # key first
    assert names(template_columns(COLUMNS, "delete", ["id"])) == ["id"]
    assert template_tsv(COLUMNS[1:3]) == "name\tprice\n"


def test_saved_template_explains_each_column(tmp_path):
    columns = COLUMNS + [column("born", t.Date())]
    save_template(tmp_path / "t.xlsx", columns)
    sheet = load_workbook(tmp_path / "t.xlsx").active
    assert [cell.value for cell in sheet[1]] == ["id", "name", "price", "active", "born"]
    assert sheet["B1"].comment.text == "VARCHAR(50), required"
    assert describe_column(columns[0]) == "INTEGER, primary key, numbered by the database"
    assert sheet["E2"].number_format == "@"  # dates stay text, so Excel can't reformat them
    assert sheet.freeze_panes == "A2"


# Writing


@pytest.fixture
def items_db(tmp_path):
    path = tmp_path / "items.sqlite3"
    conn = sqlite3.connect(path)
    conn.executescript(
        """
        CREATE TABLE items (id INTEGER PRIMARY KEY, name TEXT NOT NULL UNIQUE, price NUMERIC(10, 2));
        INSERT INTO items (name, price) VALUES ('a', 1), ('b', 2);
        """
    )
    conn.close()
    return path


ITEMS = Relation("main", "items", "table")


def writable(path):
    return Database(replace(detect(str(path)), allow_writes=True))


def rows(path):
    return sqlite3.connect(path).execute("SELECT id, name, price FROM items ORDER BY id").fetchall()


def run_import(db, operation, text, key=("id",)):
    planned = plan(parse_pasted(text), db.import_columns(ITEMS), operation, list(key))
    return db.write_rows(ITEMS, operation, list(key), planned.ready)


def test_import_columns(items_db):
    columns = {c.name: c for c in Database(detect(str(items_db))).import_columns(ITEMS)}
    assert columns["id"].generated and columns["id"].primary_key and not columns["id"].nullable
    assert not columns["name"].nullable and not columns["name"].has_default


def test_read_only_connections_refuse(items_db):
    with pytest.raises(WritesNotAllowed):
        run_import(Database(detect(str(items_db))), "insert", "name\nc\n")
    assert len(rows(items_db)) == 2


def test_insert_skips_failed_rows(items_db):
    report = run_import(writable(items_db), "insert", "name\tprice\nc\t3\na\t9\nd\t4\n")
    assert (report.inserted, [(f.line, f.error) for f in report.failures]) == (2, [(3, "UNIQUE constraint failed: items.name")])
    assert [r[1] for r in rows(items_db)] == ["a", "b", "c", "d"]


def test_update_upsert_delete(items_db):
    db = writable(items_db)
    report = run_import(db, "update", "id\tprice\n1\t10\n99\t5\n")
    assert report.updated == 1 and report.failures[0].error == "No row with id = 99."
    report = run_import(db, "upsert", "id\tname\tprice\n2\tb\t20\n50\tz\t7\n")
    assert (report.updated, report.inserted) == (1, 1)
    report = run_import(db, "delete", "id\n1\n77\n")
    assert report.deleted == 1 and report.failures[0].error == "No row with id = 77."
    assert rows(items_db) == [(2, "b", 20), (50, "z", 7)]


def test_update_by_another_key(items_db):
    report = run_import(writable(items_db), "update", "name\tprice\nb\t99\n", key=("name",))
    assert report.updated == 1 and rows(items_db)[1] == (2, "b", 99)


def test_gives_up_when_the_first_rows_all_fail(items_db, monkeypatch):
    monkeypatch.setattr(database, "GIVE_UP_AFTER", 3)
    report = run_import(writable(items_db), "insert", "name\na\nb\na\nb\nnew\n")
    assert report.stopped.startswith("Stopped after the first 3 rows all failed")
    assert report.inserted == 0 and len(report.failures) == 3


def test_duckdb_commits_row_by_row(tmp_path):
    duckdb = pytest.importorskip("duckdb")
    pytest.importorskip("duckdb_engine")
    path = tmp_path / "x.duckdb"
    conn = duckdb.connect(str(path))
    conn.execute("CREATE TABLE items (id INTEGER PRIMARY KEY, name VARCHAR NOT NULL, price DECIMAL(10, 2))")
    conn.close()
    db = writable(path)
    items = Relation("main", "items", "table")
    planned = plan(parse_pasted("id\tname\n1\ta\n1\tdup\n2\tb\n"), db.import_columns(items), "insert", ["id"])
    report = db.write_rows(items, "insert", ["id"], planned.ready)
    assert report.inserted == 2 and [f.line for f in report.failures] == [3]


# The Import tab


def run(app, scenario, size=(140, 50)):
    async def go():
        async with app.run_test(size=size) as pilot:
            await scenario(pilot)

    asyncio.run(asyncio.wait_for(go(), 25))


async def settle(pilot):
    for _ in range(5):
        await pilot.app.workers.wait_for_complete()
        await pilot.pause()


async def paste(app, pilot, text):
    from textual import events

    app.focused.post_message(events.Paste(text))  # what the terminal sends for ctrl+v
    await settle(pilot)


def text_of(app, selector):
    return str(app.import_panel.query_one(selector).content)


def test_read_only_connections_show_how_to_allow_writes(items_db):
    app = InspectorApp(detect(str(items_db)))

    async def scenario(pilot):
        await settle(pilot)
        await pilot.press("down", "f4")
        await settle(pilot)
        assert app.import_panel.relation == ITEMS  # the highlighted table
        await paste(app, pilot, "name\nc\n")
        assert "press c, then w" in text_of(app, "#import-lock")
        assert app.import_panel.query_one("#run-import", Button).disabled

    run(app, scenario)


def test_paste_preview_confirm_import(items_db):
    app = InspectorApp(replace(detect(str(items_db)), allow_writes=True))

    async def scenario(pilot):
        await settle(pilot)
        await pilot.press("down", "f4")
        await settle(pilot)
        await paste(app, pilot, "Name\tPrice\tColour\nc\t3\tred\na\t9\tblue\nd\tcheap\tgreen\n")
        assert "columns: name, price  |  ignored (no such column): Colour" in text_of(app, "#import-summary")
        assert "1 row has a problem" in text_of(app, "#import-summary")
        button = app.import_panel.query_one("#run-import", Button)
        assert str(button.label) == "Insert 2 rows" and not button.disabled

        await pilot.press("ctrl+r")
        await pilot.pause()
        assert isinstance(app.screen, ConfirmScreen) and "Insert 2 rows into items" in app.screen.question
        await pilot.press("y")
        await settle(pilot)
        assert "1 inserted  |  2 failed" in text_of(app, "#import-result")

        await pilot.click("#copy-failed")
        assert app.clipboard == (
            "Name\tPrice\tColour\terror\n"
            "a\t9\tblue\tUNIQUE constraint failed: items.name\n"
            "d\tcheap\tgreen\tprice: 'cheap' isn't a number\n"
        )
        # Fixed rows can go straight back: the error column is ignored.
        await paste(app, pilot, "Name\tPrice\terror\ne\t9\tUNIQUE constraint failed\n")
        assert "ignored (no such column): error" in text_of(app, "#import-summary")

    run(app, scenario)
    assert [r[1] for r in rows(items_db)] == ["a", "b", "c"]


def test_templates_from_the_tab(items_db, tmp_path):
    app = InspectorApp(replace(detect(str(items_db)), allow_writes=True), cwd=tmp_path)

    async def scenario(pilot):
        await settle(pilot)
        await pilot.press("down", "f4")
        await settle(pilot)
        await pilot.click("#copy-template")
        assert app.clipboard == "name\tprice\n"  # insert: no generated id
        app.import_panel.query_one("#import-op", Select).value = "update"
        await settle(pilot)
        await pilot.click("#save-template")
        await pilot.pause()

    run(app, scenario)
    sheet = load_workbook(tmp_path / "items-update-template.xlsx").active
    assert [cell.value for cell in sheet[1]] == ["id", "name", "price"]


def test_paste_fills_the_table_and_clear_empties_it(items_db):
    from textual.widgets import DataTable

    app = InspectorApp(replace(detect(str(items_db)), allow_writes=True))

    async def scenario(pilot):
        await settle(pilot)
        await pilot.press("down", "f4")
        await settle(pilot)
        preview = app.import_panel.query_one("#import-preview", DataTable)
        assert app.focused is preview
        await paste(app, pilot, "name\tprice\nc\t3\nd\t4\n")
        assert preview.row_count == 2
        await paste(app, pilot, "name\nx\n")  # a new paste replaces the rows
        assert preview.row_count == 1
        await pilot.press("ctrl+l")
        await pilot.pause()
        assert preview.row_count == 0 and app.import_panel.pasted == ""
        assert "press ctrl+v here" in text_of(app, "#import-summary")

    run(app, scenario)
